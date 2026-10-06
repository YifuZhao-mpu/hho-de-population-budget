"""Checkpoint figures of the CEC2020 real-world competition entries cited in the paper.

Main text Section 4.3 and Supplementary Sections S6.1 and S6.8 quote feasibility
counts and costs of five released competition entries (COLSHADE, EnMODE,
BPMAgES, SASS, sCMAgES) on RC01, RC06 and RC14.  They come from the per-run
files the organisers released with the entries' code; this script recomputes
every one of them.  The files are third-party data and are not redistributed:
download the repository below (the code archives are unchanged since commit
a148874a233f9fb8fa82fdae3b5739d1bba7020d) and point the script at the folder
that holds the archives.  Each archive is checked against its sha256 first.

    https://github.com/P-N-Suganthan/2020-RW-Constrained-Optimisation

Each per-run file holds 10 rows (the best solution after 0.1, 0.2, ..., 1.0 of
the problem's budget, MaxFEs = 1e5 for RC01 and RC14, 4e5 for RC06) by 25
columns (runs).  A run counts as feasible at a checkpoint when its recorded
constraint violation is 0; the organisers' violation already applies the 1e-4
equality tolerance.  EnMODE is the exception: it activates the constraints in
stages and records the violation over the constraints active so far
(EnMODE_Main.m), so before its final checkpoint a recorded 0 is not feasibility
and its per-checkpoint counts are upper bounds (a recorded violation above 0 still
means an infeasible run); its rows carry a note saying so.  BPMAgES is read from its white-box files (-wbC), the
accounting in which one approximated Jacobian costs one evaluation.  EnMODE's
archive holds a second copy of its results (Submitted code/Results_Record),
byte-different from the 'Submitted result' files used here but with the same
quoted figures on these three problems.  The organisers' ranking workbooks hold one
more EnMODE result set, read here for RC01 only.

Run:  python3 scripts/cec2020_entries.py <folder with the archives>
Writes results/analysis/cec2020_entries.csv and prints the figures.
"""
import csv
import hashlib
import io
import os
import sys
import zipfile

import numpy as np

SHA256 = {
    "CEC E-24363 codes and results.zip": "976989263d0ffe73665416ede39870b3c68809840558ea4b1b52beacc02abce5",  # EnMODE
    "CEC E-24443 codes and results.zip": "5e81bfdda722cca147f5793825abf86b21df0afcad87bfcaea023d4bc21ba5c0",  # BPMAgES
    "CEC E-24586 codes and results.zip": "2253ca84c792c592c0d7e797d7f09c6fb26ed1d971f398bf7baf719e8ecb1023",  # COLSHADE
    "GECCO COM108 codes and results.zip": "ab5b8eb34e037a5ab7cd75cadb047ab6af507768dc54e483112544e908fba978",  # sCMAgES
    "GECCO COM109 codes and results.zip": "2da1d9918ccb21561f45f2bc25cf3cdef53951453d35577a304fd7c6d29e96e7",  # SASS
    "Code used For Ranking.zip": "2f873f3ac90f67a8375730c490cb0a543eee47547181b610ee5507cbd711e291",
}
# archive, then the F and CV member paths with {p} = problem number
ENTRIES = {
    "COLSHADE": ("CEC E-24586 codes and results.zip",
                 "CEC E-24586 codes and results/Submitted Results/COLSHADE_RC{p:02d}_F.txt",
                 "CEC E-24586 codes and results/Submitted Results/COLSHADE_RC{p:02d}_CV.txt"),
    "EnMODE": ("CEC E-24363 codes and results.zip",
               "CEC E-24363 codes and results/Submitted result/EnMODE_RC{p}__F.txt",
               "CEC E-24363 codes and results/Submitted result/EnMODE_RC{p}__CV.txt"),
    "BPMAgES": ("CEC E-24443 codes and results.zip",
                "CEC E-24443 codes and results/Submitted Results/BiPopEpsMAgES_RC{p:02d}_F-wbC.txt",
                "CEC E-24443 codes and results/Submitted Results/BiPopEpsMAgES_RC{p:02d}_CV-wbC.txt"),
    "SASS": ("GECCO COM109 codes and results.zip",
             "GECCO COM109 codes and results/Submitted Result/SASS_RC{p}_F.txt",
             "GECCO COM109 codes and results/Submitted Result/SASS_RC{p}_CV.txt"),
    "sCMAgES": ("GECCO COM108 codes and results.zip",
                "GECCO COM108 codes and results/Submitted Results/sCMAgES{p}_F.txt",
                "GECCO COM108 codes and results/Submitted Results/sCMAgES{p}_CV.txt"),
}
MAXFES = {1: 100_000, 6: 400_000, 14: 100_000}
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "results", "analysis", "cec2020_entries.csv")


def open_archives(folder):
    zs = {}
    for name, digest in SHA256.items():
        path = os.path.join(folder, name)
        if not os.path.exists(path):
            raise SystemExit(f"missing {path}: download the repository named in the docstring")
        got = hashlib.sha256(open(path, "rb").read()).hexdigest()
        if got != digest:
            raise SystemExit(f"{name}: sha256 {got} differs from the archive this script was written for")
        zs[name] = zipfile.ZipFile(path)
    return zs


def runs(zs, entry, p):
    arc, f, cv = ENTRIES[entry]
    load = lambda m: np.loadtxt(io.BytesIO(zs[arc].read(m.format(p=p))))
    F, CV = load(f), load(cv)
    if F.shape != (10, 25) or CV.shape != (10, 25):
        raise SystemExit(f"{entry} RC{p:02d}: expected 10 checkpoints x 25 runs, got {F.shape}")
    return F, CV


def ranking_enmode_rc01(zs):
    import openpyxl  # only needed for this one figure
    wb = openpyxl.load_workbook(io.BytesIO(zs["Code used For Ranking.zip"].read("Code used For Ranking/EnMODE.xlsx")),
                                read_only=True, data_only=True)
    F = np.array([r[0] for r in wb["Sheet1"].iter_rows(values_only=True)], dtype=float)  # column 1 = RC01
    CV = np.array([r[0] for r in wb["Sheet2"].iter_rows(values_only=True)], dtype=float)
    return F, CV


def main():
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    zs = open_archives(sys.argv[1])
    rows = []
    for entry in ENTRIES:
        for p in (1, 6, 14):
            F, CV = runs(zs, entry, p)
            feas = CV <= 0
            per_cp = feas.sum(axis=1)
            first = next((k + 1 for k in range(10) if per_cp[k] > 0), None)
            fin = F[-1][feas[-1]]
            rows.append({
                "entry": entry, "problem": f"RC{p:02d}", "maxfes": MAXFES[p],
                "note": ("violation recorded over the constraints active so far: counts before the final "
                         "checkpoint are upper bounds" if entry == "EnMODE" else ""),
                "feasible_runs_by_checkpoint": " ".join(str(int(c)) for c in per_cp),
                "first_checkpoint_with_a_feasible_run_evals": first * MAXFES[p] // 10 if first else "",
                "feasible_runs_final": int(feas[-1].sum()),
                "best_feasible_final": f"{fin.min():.6f}" if fin.size else "",
                "median_feasible_final": f"{np.median(fin):.6f}" if fin.size else "",
                "max_feasible_final": f"{fin.max():.6f}" if fin.size else "",
                "median_feasible_checkpoint1": f"{np.median(F[0][feas[0]]):.6f}" if feas[0].any() else "",
            })
    Fr, CVr = ranking_enmode_rc01(zs)
    rows.append({"entry": "EnMODE (ranking workbook)", "problem": "RC01", "maxfes": MAXFES[1], "note": "final values only",
                 "feasible_runs_by_checkpoint": "", "first_checkpoint_with_a_feasible_run_evals": "",
                 "feasible_runs_final": int((CVr <= 0).sum()),
                 "best_feasible_final": "", "median_feasible_final": "", "max_feasible_final": "",
                 "median_feasible_checkpoint1": ""})
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    for r in rows:
        print(" | ".join(f"{k}={v}" for k, v in r.items()))
    print(f"\nwrote {os.path.normpath(OUT)}")


if __name__ == "__main__":
    main()
